"""Editable scientific manuscripts with native equations and three-line tables."""
from pathlib import Path
import json,re,argparse
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];R=ROOT/'revision_20261009';M=R/'manuscript'

def mr(text):
 r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
def seq(*nodes):return [mr(x) if isinstance(x,str) else x for x in nodes]
def script(base,sub=None,sup=None):
 kind='sSubSup' if sub is not None and sup is not None else ('sSub' if sub is not None else 'sSup')
 node=OxmlElement('m:'+kind);body=OxmlElement('m:e')
 for el in (seq(base) if isinstance(base,str) else base):body.append(el)
 node.append(body)
 for tag,val in [('sub',sub),('sup',sup)]:
  if val is not None:
   s=OxmlElement('m:'+tag);s.append(mr(val));node.append(s)
 return node
def v(x):return script(x,'t')
INLINE={'V(e_start)−V(e_end)':lambda:seq('V(',script('e','start'),')−V(',script('e','end'),')'),
 'n_hat_t':lambda:script('n̂','t'),'g_hat_t':lambda:script('ĝ','t'),
 'c_t':lambda:v('c'),'d_t':lambda:v('d'),'e_t':lambda:v('e'),
 'n_hat':lambda:mr('n̂'),'C_i':lambda:script('C','i'),
 'A_ε':lambda:script('A','ε'),'e_start':lambda:script('e','start'),'e_end':lambda:script('e','end')}
def add_prose(doc,text):
 p=doc.add_paragraph();pattern=re.compile(r'(?<!\w)('+ '|'.join(sorted(map(re.escape,INLINE),key=len,reverse=True))+r')(?!\w)')
 last=0
 for match in pattern.finditer(text):
  p.add_run(text[last:match.start()]);m=OxmlElement('m:oMath');nodes=INLINE[match.group()]()
  for node in (nodes if isinstance(nodes,list) else [nodes]):m.append(node)
  p._p.append(m);last=match.end()
 p.add_run(text[last:]);return p
def mathblock(doc,text):
 if text=='[[EQ_BATTERY]]':
  no=1;lines=[seq(v('e'),' = ',script('e','t−1'),' + Δ(η ',v('c'),' − ',v('d'),'/η),'),seq('0 ≤ ',v('c'),', ',v('d'),' ≤ 3,     1 ≤ ',v('e'),' ≤ 9,     ',script('e','0'),' = ',script('e','48'),' = 5.')]
 elif text=='[[EQ_OBJECTIVE]]':
  no=2;lines=[seq('Ĵ = Δ ',script('∑','t'),' [',script('p','t','buy'),' max(',v('ĝ'),',0) + ',script('p','t','sell'),' min(',v('ĝ'),',0)'),seq('+ κ(',v('c'),' + ',v('d'),')] + εR(c,d,e).')]
 elif text=='[[EQ_SELECTION]]':
  no=3;lines=[seq('R(c,d,e) = ',script('∑','t'),' [',script(seq('(',v('c'),'/3)'),sup='2'),' + ',script(seq('(',v('d'),'/3)'),sup='2'),' + ',script(seq('(',v('e'),'/10)'),sup='2'),'].')]
 elif text=='[[EQ_DFL]]':
  no=4;lines=[seq(script('L','DFL'),'(θ) = (1/N) ',script('∑','i'),' ',script('C','i'),'[',script('A','ε'),'(',script('l̂','i'),' − ',script('f','θ'),'(',script('x','i'),'))].')]
 else:raise ValueError(text)
 for i,nodes in enumerate(lines):
  p=doc.add_paragraph();p.alignment=1;p.paragraph_format.space_after=Pt(4)
  p.paragraph_format.keep_with_next=i<len(lines)-1;p.paragraph_format.keep_together=True
  m=OxmlElement('m:oMath')
  for node in nodes:m.append(node)
  p._p.append(m)
  if i==len(lines)-1:p.add_run(f'    ({no})')

def font(style,cn=False):
 style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
 pr=style._element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
 if rf is None:rf=OxmlElement('w:rFonts');pr.append(rf)
 for att in list(rf.attrib):
  if 'Theme' in att or 'theme' in att:del rf.attrib[att]
 for part in ('ascii','hAnsi','cs'):rf.set(qn('w:'+part),'Times New Roman')
 rf.set(qn('w:eastAsia'),'Microsoft YaHei' if cn else 'Times New Roman')

def width_weights(b):
 n=len(b['headers']);cap=b['caption']
 if cap.startswith('Table 5.'):return [1,3.8,3.4]
 if cap.startswith('Table 3.'):return [1.7,.9,1.45,1.05,3.6]
 if cap.startswith('Table S2.'):return [.8,2.6]+[1]*(n-2)
 if cap.startswith(('Table S5.','Table S7.')):return [2.2]+[1]*(n-1)
 if cap.startswith('Table S9.'):return [2]+[1]*(n-1)
 if cap.startswith('Table S4'):return [1.4,.75,1.1]+[1]*(n-3)
 if n>=6:return [1.7]+[1]*(n-1)
 return [1.8]+[1]*(n-1)

def add_table(doc,b):
 p=doc.add_paragraph(b['caption'],'Caption');p.paragraph_format.keep_with_next=True
 if b['caption'].startswith('Table 1.'):p.paragraph_format.page_break_before=True
 t=doc.add_table(rows=1,cols=len(b['headers']));t.autofit=False
 weights=width_weights(b);ws=[169*w/sum(weights) for w in weights]
 for i,w in enumerate(ws):t.columns[i].width=Mm(w)
 props=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for edge in ('top','bottom','left','right','insideH','insideV'):
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ('top','bottom') else 'nil')
  e.set(qn('w:sz'),'8');e.set(qn('w:color'),'000000');borders.append(e)
 props.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for edge,size in [('top',65),('bottom',65),('left',60),('right',60)]:
  e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(size));e.set(qn('w:type'),'dxa');margins.append(e)
 props.append(margins)
 for i,h in enumerate(b['headers']):t.rows[0].cells[i].text=h
 for vals in b['rows']:
  cs=t.add_row().cells
  for c,value in zip(cs,vals):c.text=str(value).replace('_',' ')
 for ri,row in enumerate(t.rows):
  rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
  if ri==0:rp.append(OxmlElement('w:tblHeader'))
  for ci,cell in enumerate(row.cells):
   cell.width=Mm(ws[ci]);cell.vertical_alignment=1
   if ri==0:
    cb=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');cb.append(e);cell._tc.get_or_add_tcPr().append(cb)
   for p in cell.paragraphs:
    p.paragraph_format.space_before=Pt(1);p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
    if ri==0 or (b['caption'].startswith(('Table 1.','Table 2.','Table 4.')) and ri<len(t.rows)-1):p.paragraph_format.keep_with_next=True
    p.alignment=0 if ci==0 or b['caption'].startswith(('Table 3.','Table 5.')) else 1
    for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
 p=doc.add_paragraph(b['note']);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(10)
 for run in p.runs:run.font.size=Pt(9)

def build(which):
 names={'article':('article_blocks.json','EPSR_Revised_Manuscript.docx'),'supplement':('supplement_blocks.json','EPSR_Supplementary_Material.docx'),'response':('response_blocks.json','逐条修改说明.docx')}
 src,filename=names[which];data=json.loads((M/src).read_text(encoding='utf-8'))
 doc=Document();s=doc.sections[0]
 # Preserve the supplied manuscript's A4 scientific-paper page size.
 s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(20);s.left_margin=s.right_margin=Mm(20.5)
 for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
  st=doc.styles[name];font(st,which=='response');st.font.size=Pt(11);st.paragraph_format.space_after=Pt(6)
 doc.styles['Normal'].paragraph_format.line_spacing=1.1
 doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
 for name in ['Heading 1','Heading 2']:
  doc.styles[name].font.size=Pt(12 if name=='Heading 1' else 11);doc.styles[name].font.bold=True
  doc.styles[name].paragraph_format.keep_with_next=True;doc.styles[name].paragraph_format.space_before=Pt(9)
 doc.styles['Caption'].font.size=Pt(10)
 for b in data['blocks']:
  kind=b['type']
  if kind in ('title','heading'):
   style='Title' if kind=='title' else ('Heading 2' if re.match(r'^\d\.\d',b['text']) else 'Heading 1')
   doc.add_paragraph(b['text'],style)
  elif kind=='table':add_table(doc,b)
  elif kind=='figure':
   imagepath=Path(b['path'].replace('.preview.png','.png'));assert imagepath.exists()
   p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(imagepath),width=Mm(169))
   p=doc.add_paragraph(b['caption'],'Caption');p.paragraph_format.keep_with_next=True
   p=doc.add_paragraph(b['note'])
   for run in p.runs:run.font.size=Pt(9)
  elif kind=='equation':mathblock(doc,b['text'])
  else:
   p=add_prose(doc,b['text']) if kind=='paragraph' else doc.add_paragraph(b['text'])
   if kind=='reference':
    p.paragraph_format.space_after=Pt(5)
    for run in p.runs:run.font.size=Pt(9.5)
 footer=s.footer.paragraphs[0];footer.alignment=2
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
 for border in doc._element.xpath('.//w:pBdr'):border.getparent().remove(border)
 for border in doc.styles._element.xpath('.//w:pBdr'):border.getparent().remove(border)
 doc.core_properties.title=data['title'];doc.core_properties.subject='Substantive revision for Electric Power Systems Research';doc.core_properties.author=''
 out=R/filename;doc.save(out)
 print(json.dumps({'file':str(out),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'math_objects':len(doc._element.xpath('.//m:oMath')),'bytes':out.stat().st_size},ensure_ascii=False))
if __name__=='__main__':
 a=argparse.ArgumentParser();a.add_argument('--document',choices=['article','supplement','response','all'],default='all');args=a.parse_args()
 for which in (['article','supplement','response'] if args.document=='all' else [args.document]):build(which)
