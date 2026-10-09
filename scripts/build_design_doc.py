from pathlib import Path
import json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.73); sec.bottom_margin=Inches(.68)
sec.left_margin=sec.right_margin=Inches(.79)
sec.header_distance=Inches(.3); sec.footer_distance=Inches(.3)
for name,size in [('Normal',11),('Title',20),('Subtitle',11),('Heading 1',14),('Heading 2',12)]:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun' if name=='Normal' else 'Microsoft YaHei')
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        st.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
    for b in st.element.xpath('.//w:pBdr'):b.getparent().remove(b)
    if name in ['Title','Subtitle']:st.font.italic=False
    st.paragraph_format.line_spacing=1.23
    st.paragraph_format.space_after=Pt(7)
    if name.startswith('Heading'): st.paragraph_format.space_before=Pt(13)
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
header=sec.header.paragraphs[0];header.text='住宅光储决策学习研究方案   2026年10月5日'
header.style=doc.styles['Normal']
for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run(); fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)

def link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'1F4E79');rp.append(co)
    r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    props=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for vals, row in [(headers,t.rows[0])]+[(vals,t.add_row()) for vals in rows]:
        trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        if row==t.rows[0]: pass
        for i,v in enumerate(vals):
            c=row.cells[i];c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for e in ['top','bottom','left','right']:
                x=OxmlElement('w:'+e);x.set(qn('w:w'),'95');x.set(qn('w:type'),'dxa');m.append(x)
            pr.append(m)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
            p.add_run(str(v))
            for run in p.runs:run.font.size=Pt(10)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for c in t.rows[0].cells:
        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');c._tc.get_or_add_tcPr().append(shade)
        for r in c.paragraphs[0].runs:r.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

text=(ROOT/'research/report_content.txt').read_text(encoding='utf-8')
for line in text.splitlines():
    if not line.strip():continue
    if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('基于公共') or line.startswith('检索与设计'):doc.add_paragraph(line,'Subtitle')
    elif ' https://' in line and len(line)<180:
        label,url=line.split(' https://',1);p=doc.add_paragraph();link(p,label,'https://'+url)
    else:doc.add_paragraph(line)

doc.add_paragraph('14 已完成的本机核验','Heading 1')
panel=json.loads((ROOT/'datasets/panel_manifest.json').read_text(encoding='utf-8'))
smoke=json.loads((ROOT/'results/smoke_dispatch_real_sample.json').read_text(encoding='utf-8'))
table(['项目','本次状态','边界'],[
 ['文献','11篇核心 其中7篇全文','另有1篇部分正文和3篇摘要导言'],
 ['Ausgrid原包','完整下载 全部ZIP成员CRC通过','镜像来源 已核对政府目录许可'],
 ['数据面板','300户 × 1096日 × 48槽位 × 3通道','已转为D盘处理副本 保留质量与缺失掩码'],
 ['严格有效数据',f"GC/GG {panel['valid_gc_gg_household_days']:,} 户日\n含CL质量要求 {panel['valid_feeder_household_days']:,} 户日",'已统一排除客户2 尚未选定最终住宅面板'],
 ['真实户日LP','能量平衡残差约2.22×10⁻¹⁶\n终端能量5 kWh 无同时充放电','仅实现检查 不代表经济或网络研究结果'],
 ['正式算法与潮流','尚未运行','可微层 梯度 DFL AC潮流及全年结果均待执行']
],[1.02,2.68,2.99])
doc.add_paragraph('面板转换不进行插值或训练，不用测试数据拟合任何统计量。全过程保留原始源时钟；不能把该面板当作已经完成夏令时校正的UTC数据。原始数据保持只读，失败下载残片没有覆盖成功版本。')

doc.add_paragraph('15 核心文献与证据入口','Heading 1')
papers=json.loads((ROOT/'research/literature_decisions.json').read_text(encoding='utf-8'))['papers']
access={'full_text_author_version':'作者全文','full_text_publisher_pdf_in_institutional_repository':'机构存档出版全文','full_text_preprint':'预印本全文 尚未经同行评审','abstract_and_publisher_introduction_only':'仅摘要与导言','abstract_plus_indexed_fulltext_sections_4_1_4_2':'摘要及索引正文片段'}
for x in papers:
    p=doc.add_paragraph()
    p.paragraph_format.keep_with_next=True
    r=p.add_run(f"〔{x['id']}〕{'; '.join(x['authors'])}. {x['title']}. {x['venue']}, {x['year']}. DOI {x['doi']}.")
    r.font.size=Pt(10)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10)
    p.add_run('核验层级：'+access.get(x['access'],x['access'])+'。 ')
    link(p,'原始来源',x['original_url'])
    if x.get('fulltext_url'):p.add_run('  ');link(p,'全文入口',x['fulltext_url'])
    for r in p.runs:r.font.size=Pt(10)

doc.add_paragraph('16 数据与网络资料入口','Heading 1')
refs=[
 ('Ausgrid 官方政府目录与原始许可','https://data.gov.au/data/dataset/5ab48b70-5e99-47d3-9193-5c34a2676d93'),
 ('Ausgrid 镜像来源与保存说明','https://github.com/pierre-haessig/ausgrid-solar-data'),
 ('Ausgrid 完整数据下载','https://pierreh.eu/downloads/Ausgrid_solar_home_data.zip'),
 ('OPSD Household Data 2020年4月15日版','https://data.open-power-system-data.org/household_data/2020-04-15/'),
 ('OPSD 元数据与许可','https://data.open-power-system-data.org/household_data/2020-04-15/datapackage.json'),
 ('IEEE测试馈线工作组仓库','https://github.com/ieee-pes-amps/dtf-dev'),
 ('IEEE仓库许可正文 当前标为TBD','https://raw.githubusercontent.com/ieee-pes-amps/dtf-dev/master/LICENSE.md'),
 ('IEEE European LV馈线工作组说明','https://ewh.ieee.org/soc/pes/dsacom/testfeeders/Minutes07282015.pdf'),
 ('EPRI OpenDSS 官方文档','https://opendss.epri.com/'),
 ('SimBench数据入口','https://simbench.de/en/download/datasets/'),
 ('SimBench许可','https://raw.githubusercontent.com/e2nIEE/simbench/develop/LICENSE'),
 ('SimBench模型说明与三相对称假设','https://simbench.de/wp-content/uploads/2020/01/simbench_documentation_en_1.0.0.pdf'),
 ('运行边界近邻 Operating Envelopes under Probabilistic Electricity Demand and Solar Generation Forecasts','https://arxiv.org/abs/2207.09818'),
 ('运行边界近邻 Coordinated Dynamic Operating Envelopes for Unlocking Additional Flexibility at Grid Edge 预印本','https://arxiv.org/abs/2604.17081')]
for label,url in refs:
    p=doc.add_paragraph();link(p,label,url);p.paragraph_format.space_after=Pt(6)
doc.core_properties.title='住宅光储决策学习的配电网影响研究方案'
doc.core_properties.subject='公共数据 近邻文献 受控仿真 可执行设计'
doc.core_properties.author=''
output=ROOT/'住宅光储决策学习研究方案.docx'
doc.save(output)
print(output)
print('paragraphs',len(doc.paragraphs),'characters',len(text),'tables',len(doc.tables))
